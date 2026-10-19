"""Documents sent on Telegram: text out of PDF, DOCX, TXT, MD, CSV and JSON.

The file goes to the vault's ``05-anexos`` folder with a note next to it that
holds the extracted text (so the search index and recall find it); the caption
goes to the router with the first part of the text in the context block, and
``document_read`` reads the rest.
"""

from __future__ import annotations

import io
import json
from dataclasses import dataclass
from pathlib import PurePath

from uva.i18n import t

TEXT_TYPES = (".txt", ".md", ".csv", ".json")
SUPPORTED = (".pdf", ".docx", *TEXT_TYPES)
# What the LLM reads first, and what document_read returns per call.
READ_CHARS = 8000


class DocumentError(ValueError):
    """A document that can't be read (shown to the user)."""


@dataclass(frozen=True, slots=True)
class Extracted:
    text: str
    truncated: bool  # longer than DOCUMENT_MAX_CHARS
    pages: int | None = None  # PDFs


def extension(filename: str) -> str:
    return PurePath(filename or "").suffix.lower()


def supported(filename: str) -> bool:
    return extension(filename) in SUPPORTED


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _pdf(data: bytes) -> tuple[str, int]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:  # broken or encrypted file
        raise DocumentError(t("documents.unreadable", error=type(exc).__name__)) from exc
    return "\n\n".join(p.strip() for p in pages if p.strip()), len(pages)


def _docx(data: bytes) -> str:
    from docx import Document

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentError(t("documents.unreadable", error=type(exc).__name__)) from exc
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def extract(filename: str, data: bytes, *, max_chars: int) -> Extracted:
    """The document's text, clipped to ``max_chars``. Raises DocumentError."""
    ext = extension(filename)
    pages = None
    if ext == ".pdf":
        text, pages = _pdf(data)
    elif ext == ".docx":
        text = _docx(data)
    elif ext in TEXT_TYPES:
        text = _decode(data)
        if ext == ".json":
            try:  # pretty-printed reads better
                text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
            except ValueError:
                pass
    else:
        raise DocumentError(t("documents.unsupported", ext=ext or "?", types=", ".join(SUPPORTED)))
    text = text.replace("\r\n", "\n").strip()
    if not text:
        raise DocumentError(t("documents.empty"))
    truncated = len(text) > max_chars
    return Extracted(text[:max_chars], truncated, pages)


def note_body(attachment: str, extracted: Extracted) -> str:
    """The note kept next to the file: a link to it, then the text."""
    return f"![[{attachment}]]\n\n{extracted.text}"


def text_of(note_body_text: str) -> str:
    """The extracted text of a document note (without the link to the file)."""
    text = note_body_text.lstrip()
    if text.startswith("# "):  # the title heading the vault writes
        text = text.partition("\n")[2].lstrip()
    head, sep, rest = text.partition("\n\n")
    return rest if sep and head.startswith("![[") else text


def window(text: str, offset: int, size: int = READ_CHARS) -> tuple[str, int | None]:
    """``size`` chars from ``offset``; the next offset, or None at the end."""
    offset = max(0, offset)
    chunk = text[offset : offset + size]
    end = offset + len(chunk)
    return chunk, (end if end < len(text) else None)
